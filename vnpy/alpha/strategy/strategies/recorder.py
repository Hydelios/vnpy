import os
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, date, time
from pathlib import Path
from typing import Dict, List, Any
from vnpy.trader.constant import Direction, Offset, Status
from vnpy.trader.object import OrderData, TradeData
import json
import polars as pl


def _normalize_columns(columns: List[Any]) -> List[str]:
    normalized: List[str] = []
    seen: set[str] = set()

    for col in columns:
        if col is None:
            continue

        name = str(col).strip()
        if not name or name.lower() == "none":
            continue

        if name in seen:
            continue

        seen.add(name)
        normalized.append(name)

    return normalized


def _ensure_columns(df: pl.DataFrame, columns: List[str]) -> pl.DataFrame:
    columns = _normalize_columns(columns)
    missing = [col for col in columns if col not in df.columns]
    if missing:
        df = df.with_columns([pl.lit(None).alias(col) for col in missing])
    return df.select(columns)


def _empty_df(columns: List[str]) -> pl.DataFrame:
    columns = _normalize_columns(columns)
    return pl.DataFrame({col: [] for col in columns})


def _write_excel_sheets(path: Path, sheets: Dict[str, pl.DataFrame]) -> None:
    writer_cls = getattr(pl, "ExcelWriter", None)
    if writer_cls is not None:
        with writer_cls(path) as writer:
            for name, df in sheets.items():
                try:
                    df.write_excel(writer, worksheet=name)
                except TypeError:
                    df.write_excel(writer, sheet_name=name)
        return

    try:
        from openpyxl import Workbook
    except Exception as exc:
        raise RuntimeError("未检测到可用的 Excel 写入器，请安装 openpyxl 或升级 polars。") from exc

    wb = Workbook()
    first = True
    for name, df in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False

        columns = list(df.columns)
        if columns:
            ws.append(columns)
            for row in df.iter_rows():
                ws.append(list(row))

    wb.save(path)


ICBM_REBALANCE_COLUMNS = [
    "*证券代码",
    "*交易方向",
    "指令价格",
    "*指令数量",
    "*指令金额",
    "*交易市场",
    "指令价格相对昨收盘增幅(%)",
]

O32_REBALANCE_COLUMNS = [
    "证券代码",
    "委托方向",
    "指令数量",
    "指令价格",
    "价格模式",
    "交易市场内部编号",
    "当前指令市值/净值(%)",
    "目标市值/净值(%)",
    "账户编号/序号",
    "账户名称",
    "组合编号",
    "组合名称",
    "投资类型",
    "券商",
    "指令金额",
    "预置开始日期",
    "预置结束日期",
    "预置开始时间",
    "预置结束时间",
    "触发模式",
    "触发参数",
    "当日数量",
    "比较方向",
    "指数下限",
    "指数上限",
    "目标数量",
    "备注",
    "持仓数量比例",
    "操作级别",
    "转入组合编号",
    "转入投资类型",
]

class DailyTradeRecorder:
    """按天记录成交和持仓信息"""
    
    def __init__(self, output_dir: str = "trade_records"):
        self.output_dir = Path(output_dir)
        self.output_dir: Path = self.output_dir.joinpath("results")
        
        self.output_dir.mkdir(parents=True, exist_ok=True)
        # 存储每日的数据
        self.daily_trades: Dict[date, List[Dict]] = {}
        self.daily_positions: Dict[date, Dict[str, float]] = {}
        self.current_positions: Dict[str, float] = {}  # 当前持仓 {symbol: position}
        
        # 记录每个标的的开仓均价
        self.avg_prices: Dict[str, Dict[str, float]] = {}  # {symbol: {"long": price, "short": price}}
        
    def record_trade(self, trade: TradeData, cash: float) -> None:
        """记录单笔成交"""
        trade_date = trade.datetime.date()
        
        if trade_date not in self.daily_trades:
            self.daily_trades[trade_date] = []
            self.daily_positions[trade_date] = self.current_positions.copy()
        
        # 计算持仓变化
        symbol = trade.symbol
        size = 1  # 合约乘数，可根据实际情况调整
        
        if trade.direction == Direction.LONG:
            position_change = trade.volume * size if trade.offset == Offset.OPEN else -trade.volume * size
        else:  # SHORT
            position_change = -trade.volume * size if trade.offset == Offset.OPEN else trade.volume * size
        
        # 更新当前持仓
        current_position = self.current_positions.get(symbol, 0)
        new_position = current_position + position_change
        
        # 更新持仓
        if abs(new_position) < 1e-6:  # 接近0
            self.current_positions.pop(symbol, None)
            self.avg_prices.pop(symbol, None)
        else:
            self.current_positions[symbol] = new_position
            self._update_avg_price(symbol, trade, current_position, new_position)
        
        # 记录成交信息
        trade_record = {
            '交易时间': trade.datetime.strftime('%Y-%m-%d %H:%M:%S.%f')[:-3],
            '合约代码': trade.symbol,
            '交易所': str(trade.exchange),
            '买卖方向': '买入' if trade.direction == Direction.LONG else '卖出',
            '开平标志': self._get_offset_text(trade.offset),
            '成交价格': trade.price,
            '成交数量': trade.volume,
            '成交金额': trade.price * trade.volume * size,
            '成交编号': trade.tradeid,
            '订单编号': trade.orderid,
            '持仓变化': position_change,
            '持仓后数量': new_position,
            '现金余额': cash,
            '手续费': 0,  # 可根据需要计算
            '策略名称': getattr(trade, 'strategy_name', 'Unknown'),
        }
        
        self.daily_trades[trade_date].append(trade_record)
        
        # 更新该日持仓快照
        self.daily_positions[trade_date] = self.current_positions.copy()
        
        # 可选：实时保存到文件
        if len(self.daily_trades[trade_date]) % 10 == 0:
            self.save_daily_data(trade_date)
    
    def _get_offset_text(self, offset: Offset) -> str:
        """获取开平仓文本"""
        if offset == Offset.OPEN:
            return "开仓"
        elif offset == Offset.CLOSE:
            return "平仓"
        elif offset == Offset.CLOSETODAY:
            return "平今仓"
        elif offset == Offset.CLOSEYESTERDAY:
            return "平昨仓"
        else:
            return str(offset)
    
    def _update_avg_price(self, symbol: str, trade: TradeData, 
                         old_position: float, new_position: float) -> None:
        """更新持仓均价"""
        if symbol not in self.avg_prices:
            self.avg_prices[symbol] = {"long": 0.0, "short": 0.0}
        
        avg_price = self.avg_prices[symbol]
        size = 1  # 合约乘数
        
        if trade.direction == Direction.LONG and trade.offset == Offset.OPEN:
            # 开多仓
            if old_position >= 0:
                # 已有持仓或空仓
                total_value = avg_price["long"] * abs(old_position) + trade.price * trade.volume * size
                avg_price["long"] = total_value / abs(new_position) if new_position != 0 else 0
            else:
                # 空头持仓，先平空再开多
                avg_price["long"] = trade.price
        elif trade.direction == Direction.SHORT and trade.offset == Offset.OPEN:
            # 开空仓
            if old_position <= 0:
                total_value = avg_price["short"] * abs(old_position) + trade.price * trade.volume * size
                avg_price["short"] = total_value / abs(new_position) if new_position != 0 else 0
            else:
                avg_price["short"] = trade.price
    
    def save_daily_data(self, target_date: date = None) -> None:
        """保存指定日期的数据，如不指定则保存所有数据"""
        if target_date:
            dates_to_save = [target_date]
        else:
            dates_to_save = list(self.daily_trades.keys())
        
        for trade_date in dates_to_save:
            if trade_date not in self.daily_trades:
                continue
                
            # 1. 保存成交明细
            trades = self.daily_trades[trade_date]
            if trades:
                trades_df = pl.DataFrame(trades)
                trades_file = self.output_dir / f"trades_{trade_date.strftime('%Y%m%d')}.xlsx"

                sheets = {"成交明细": trades_df}

                # 当日持仓汇总
                positions = self.daily_positions.get(trade_date, {})
                if positions:
                    position_data = []
                    for symbol, pos in positions.items():
                        avg_long = self.avg_prices.get(symbol, {}).get("long", 0)
                        avg_short = self.avg_prices.get(symbol, {}).get("short", 0)

                        position_data.append({
                            '合约代码': symbol,
                            '持仓数量': pos,
                            '持仓方向': '多头' if pos > 0 else '空头',
                            '持仓市值': abs(pos) * (avg_long if pos > 0 else avg_short),
                            '开仓均价(多头)': avg_long if pos > 0 else 0,
                            '开仓均价(空头)': avg_short if pos < 0 else 0,
                        })

                    positions_df = pl.DataFrame(position_data)
                    sheets["持仓汇总"] = positions_df

                # 当日统计
                if trades_df.height:
                    buy_count = trades_df.filter(pl.col("买卖方向") == "买入").height
                    sell_count = trades_df.filter(pl.col("买卖方向") == "卖出").height
                    open_count = trades_df.filter(pl.col("开平标志").str.contains("开仓")).height
                    close_count = trades_df.filter(pl.col("开平标志").str.contains("平仓")).height

                    stats_df = pl.DataFrame({
                        "统计项目": [
                            "总成交笔数",
                            "总成交量",
                            "总成交额",
                            "买入次数",
                            "卖出次数",
                            "开仓次数",
                            "平仓次数",
                            "最大单笔成交",
                            "平均成交价",
                        ],
                        "数值": [
                            trades_df.height,
                            trades_df["成交数量"].sum(),
                            trades_df["成交金额"].sum(),
                            buy_count,
                            sell_count,
                            open_count,
                            close_count,
                            trades_df["成交金额"].max(),
                            trades_df["成交价格"].mean(),
                        ],
                    })
                    sheets["当日统计"] = stats_df

                _write_excel_sheets(trades_file, sheets)
                print(f"已保存 {trade_date} 的成交数据到: {trades_file}")
    
    def save_all_data(self) -> None:
        """保存所有日期的数据"""
        self.save_daily_data()
        
        # # 创建汇总文件
        # summary_file = self.output_dir / "trading_summary.xlsx"
        
        # with pd.ExcelWriter(summary_file, engine='openpyxl') as writer:
        #     # 汇总所有日期的成交
        #     all_trades = []
        #     for date_val, trades in self.daily_trades.items():
        #         for trade in trades:
        #             trade['日期'] = date_val
        #             all_trades.append(trade)
            
        #     if all_trades:
        #         all_trades_df = pd.DataFrame(all_trades)
        #         all_trades_df.to_excel(writer, sheet_name='所有成交', index=False)
            
        #     # 按日期汇总统计
        #     daily_stats = []
        #     for date_val in sorted(self.daily_trades.keys()):
        #         trades = self.daily_trades[date_val]
        #         if trades:
        #             trades_df = pd.DataFrame(trades)
        #             positions = self.daily_positions.get(date_val, {})
                    
        #             daily_stats.append({
        #                 '日期': date_val,
        #                 '成交笔数': len(trades_df),
        #                 '总成交量': trades_df['成交数量'].sum(),
        #                 '总成交额': trades_df['成交金额'].sum(),
        #                 '持仓品种数': len(positions),
        #                 '总持仓量': sum(abs(p) for p in positions.values()),
        #                 '买入次数': len(trades_df[trades_df['买卖方向'] == '买入']),
        #                 '卖出次数': len(trades_df[trades_df['买卖方向'] == '卖出']),
        #             })
            
        #     if daily_stats:
        #         daily_stats_df = pd.DataFrame(daily_stats)
        #         daily_stats_df.to_excel(writer, sheet_name='每日统计', index=False)
        
        # return {
        #     'summary': str(summary_file),
        #     'daily_files': [str(self.output_dir / f"trades_{d.strftime('%Y%m%d')}.xlsx") 
        #                   for d in self.daily_trades.keys()]
        # }


class DailyRebalanceRecorder:
    """按天记录调仓/持仓清单"""

    def __init__(
        self,
        output_dir: str = "rebalance_records/results",
        save_empty: bool = True,
        account_id: str | None = None,
        asset_unit: str | None = None,
        direction_map: Dict[str, str] | None = None,
        market_map: Dict[str, str] | None = None,
        position_template_path: str | None = None,
        rebalance_template: str | None = None,
    ):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.save_empty: bool = save_empty
        self.account_id: str | None = account_id
        self.asset_unit: str | None = asset_unit
        self.direction_map: Dict[str, str] = direction_map or {"buy": "1", "sell": "2"}
        self.market_map: Dict[str, str] = market_map or {"SSE": "1", "SZSE": "2"}
        self.wind_map: Dict[str, str] = {"SSE": "SH", "SZSE": "SZ"}
        self.rebalance_template: str = (rebalance_template or "icbm").lower()
        if self.rebalance_template not in {"icbm", "o32"}:
            raise ValueError(f"未知调仓模板类型: {self.rebalance_template}")

        if position_template_path:
            self.position_template_path = Path(position_template_path)
        else:
            repo_root = Path(__file__).resolve().parents[5]
            self.position_template_path = repo_root / "playground/alpha_research/research_factortemplate/posinfo/持仓明细报表.xlsx"

        self.position_columns: List[str] | None = None
        self.daily_rebalances: Dict[date, List[Dict[str, Any]]] = {}
        self.daily_positions: Dict[date, List[Dict[str, Any]]] = {}

    def _parse_vt_symbol(self, vt_symbol: str) -> tuple[str, str]:
        if "." in vt_symbol:
            symbol, exchange = vt_symbol.split(".", 1)
        else:
            symbol, exchange = vt_symbol, ""
        return symbol, exchange

    def _to_wind_code(self, vt_symbol: str) -> str:
        symbol, exchange = self._parse_vt_symbol(vt_symbol)
        suffix = self.wind_map.get(exchange)
        return f"{symbol}.{suffix}" if suffix else symbol

    def _to_market_code(self, vt_symbol: str) -> str:
        _, exchange = self._parse_vt_symbol(vt_symbol)
        return self.market_map.get(exchange, "")

    def _get_rebalance_columns(self) -> List[str]:
        if self.rebalance_template == "o32":
            return O32_REBALANCE_COLUMNS
        return ICBM_REBALANCE_COLUMNS

    def _load_position_columns(self) -> List[str]:
        if self.position_columns is not None:
            return self.position_columns

        default_columns = [
            "持仓日期",
            "证券代码",
            "万得代码",
            "持仓数量",
            "持仓市值（人民币）",
            "资金账号",
        ]

        def _extract_row_columns(row: ET.Element, ns: Dict[str, str], shared_strings: List[str]) -> List[str]:
            cells = {}
            for c in row.findall("a:c", ns):
                ref = c.get("r", "")
                col_ref = "".join([ch for ch in ref if ch.isalpha()])
                if not col_ref:
                    continue

                col_idx = 0
                for ch in col_ref:
                    col_idx = col_idx * 26 + (ord(ch.upper()) - ord("A") + 1)

                v = c.find("a:v", ns)
                if v is None:
                    value = None
                else:
                    value = v.text
                    if c.get("t") == "s":
                        try:
                            value = shared_strings[int(value)]
                        except Exception:
                            pass
                cells[col_idx] = value

            ordered_values = [cells[i] for i in sorted(cells.keys())]
            return _normalize_columns(ordered_values)

        try:
            if not self.position_template_path.exists():
                self.position_columns = default_columns
                return self.position_columns

            with zipfile.ZipFile(self.position_template_path) as zf:
                shared_strings = []
                if "xl/sharedStrings.xml" in zf.namelist():
                    data = zf.read("xl/sharedStrings.xml")
                    root = ET.fromstring(data)
                    ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    for si in root.findall("a:si", ns):
                        parts = []
                        for t in si.findall(".//a:t", ns):
                            parts.append(t.text or "")
                        shared_strings.append("".join(parts))

                sheet_data = zf.read("xl/worksheets/sheet1.xml")
                root = ET.fromstring(sheet_data)
                ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                rows = root.findall("a:sheetData/a:row", ns)

                columns: List[str] = []
                for row in rows[:80]:
                    row_columns = _extract_row_columns(row, ns, shared_strings)
                    if not row_columns:
                        continue

                    has_code = "证券代码" in row_columns or "万得代码" in row_columns
                    has_qty = "持仓数量" in row_columns or "数量" in row_columns
                    if has_code and has_qty:
                        columns = row_columns
                        break

                self.position_columns = columns or default_columns
        except Exception:
            self.position_columns = default_columns

        return self.position_columns

    def record_rebalance(
        self,
        target_date: date,
        pos_data: Dict[str, float],
        target_data: Dict[str, float],
        close_prices: Dict[str, float] | None = None,
        sizes: Dict[str, float] | None = None,
        position_date: date | None = None,
        strategy_name: str | None = None,
    ) -> None:
        """记录单日调仓清单"""
        close_prices = close_prices or {}
        sizes = sizes or {}
        position_date = position_date or target_date

        pos_symbols = {k for k, v in pos_data.items() if abs(v) > 1e-12}
        target_symbols = {k for k, v in target_data.items() if abs(v) > 1e-12}
        all_symbols = pos_symbols | target_symbols
        records: List[Dict[str, Any]] = []

        for vt_symbol in sorted(all_symbols):
            current_pos = pos_data.get(vt_symbol, 0)
            target_pos = target_data.get(vt_symbol, 0)
            diff = target_pos - current_pos

            if abs(diff) < 1e-12:
                continue

            direction_code = self.direction_map["buy"] if diff > 0 else self.direction_map["sell"]
            close_price = close_prices.get(vt_symbol, None)
            size = sizes.get(vt_symbol, 1)
            amount = None
            if close_price is not None:
                amount = abs(diff) * close_price * size
            symbol, _ = self._parse_vt_symbol(vt_symbol)
            if self.rebalance_template == "o32":
                record = {
                    "证券代码": symbol,
                    "委托方向": direction_code,
                    "指令数量": None,
                    "指令价格": 0,
                    "价格模式": "4",
                    "交易市场内部编号": self._to_market_code(vt_symbol),
                    "当前指令市值/净值(%)": None,
                    "目标市值/净值(%)": None,
                    "账户编号/序号": None,
                    "账户名称": None,
                    "组合编号": None,
                    "组合名称": None,
                    "投资类型": "1",
                    "券商": None,
                    "指令金额": amount,
                    "预置开始日期": None,
                    "预置结束日期": None,
                    "预置开始时间": None,
                    "预置结束时间": None,
                    "触发模式": None,
                    "触发参数": None,
                    "当日数量": None,
                    "比较方向": None,
                    "指数下限": None,
                    "指数上限": None,
                    "目标数量": None,
                    "备注": None,
                    "持仓数量比例": None,
                    "操作级别": None,
                    "转入组合编号": None,
                    "转入投资类型": None,
                }
                if diff < 0:
                    record["指令数量"] = abs(diff)
                    record["指令价格"] = 0
                    record["指令金额"] = None
            else:
                record = {
                    "*证券代码": symbol,
                    "*交易方向": direction_code,
                    "指令价格": 0,
                    "*指令数量": None,
                    "*指令金额": amount,
                    "*交易市场": self._to_market_code(vt_symbol),
                    "指令价格相对昨收盘增幅(%)": None,
                }
            records.append(record)

        position_records: List[Dict[str, Any]] = []

        final_symbols = {k for k, v in target_data.items() if abs(v) > 1e-12}
        for vt_symbol in sorted(final_symbols):
            current_pos = target_data.get(vt_symbol, 0)
            if abs(current_pos) < 1e-12:
                continue

            close_price = close_prices.get(vt_symbol, None)
            size = sizes.get(vt_symbol, 1)
            market_value = None
            if close_price is not None:
                market_value = abs(current_pos) * close_price * size

            record = {
                "持仓日期": position_date,
                "证券代码": self._parse_vt_symbol(vt_symbol)[0],
                "万得代码": self._to_wind_code(vt_symbol),
                "持仓数量": current_pos,
                "持仓市值（人民币）": market_value,
                "资金账号": self.account_id,
            }
            position_records.append(record)

        self.daily_positions[target_date] = position_records
        self.daily_rebalances[target_date] = records
        self.save_daily_data(target_date, position_date)

    def save_daily_data(self, target_date: date, position_date: date) -> None:
        """保存指定日期的调仓清单"""
        records = self.daily_rebalances.get(target_date, [])
        positions = self.daily_positions.get(target_date, [])

        if not records and not positions and not self.save_empty:
            return

        rebalance_columns = self._get_rebalance_columns()
        position_columns = self._load_position_columns()
        if records:
            rebalance_df = pl.DataFrame(records)
            rebalance_df = _ensure_columns(rebalance_df, rebalance_columns)
        else:
            rebalance_df = _empty_df(rebalance_columns)

        if positions:
            positions_df = pl.DataFrame(positions)
            positions_df = _ensure_columns(positions_df, position_columns)
        else:
            positions_df = _empty_df(position_columns)

        if self.asset_unit and not positions_df.is_empty():
            asset_cols = [col for col in position_columns if "资产单元" in col]
            for col in asset_cols:
                positions_df = positions_df.with_columns(pl.lit(self.asset_unit).alias(col))

        rebalance_path = self.output_dir / f"rebalance_{target_date.strftime('%Y%m%d')}.xlsx"
        rebalance_buy_path = self.output_dir / f"rebalance_buy_{target_date.strftime('%Y%m%d')}.xlsx"
        rebalance_sell_path = self.output_dir / f"rebalance_sell_{target_date.strftime('%Y%m%d')}.xlsx"
        position_path = self.output_dir / f"position_{position_date.strftime('%Y%m%d')}.xlsx"

        if self.save_empty or records:
            _write_excel_sheets(rebalance_path, {"调仓清单": rebalance_df})
            print(f"已保存 {target_date} 的调仓清单到: {rebalance_path}")

            direction_col = "委托方向" if self.rebalance_template == "o32" else "*交易方向"
            buy_code = self.direction_map.get("buy", "1")
            sell_code = self.direction_map.get("sell", "2")

            if direction_col in rebalance_df.columns:
                buy_df = rebalance_df.filter(pl.col(direction_col).cast(pl.Utf8) == str(buy_code))
                sell_df = rebalance_df.filter(pl.col(direction_col).cast(pl.Utf8) == str(sell_code))
            else:
                buy_df = _empty_df(rebalance_columns)
                sell_df = _empty_df(rebalance_columns)

            if self.save_empty or not buy_df.is_empty():
                _write_excel_sheets(rebalance_buy_path, {"调仓清单": buy_df})
                print(f"已保存 {target_date} 的买入清单到: {rebalance_buy_path}")
            if self.save_empty or not sell_df.is_empty():
                _write_excel_sheets(rebalance_sell_path, {"调仓清单": sell_df})
                print(f"已保存 {target_date} 的卖出清单到: {rebalance_sell_path}")

        if self.save_empty or positions:
            _write_excel_sheets(position_path, {"持仓清单": positions_df})
            print(f"已保存 {target_date} 的持仓清单到: {position_path}")
